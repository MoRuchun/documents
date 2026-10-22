"""Link numeric citations to Word bibliography bookmarks and SEQ/REF fields."""
from copy import deepcopy
from pathlib import Path
import argparse
import json
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

STYLE = '文献引用'
PREFIX = 'Bib_'
SEQUENCE = 'BibliographyItem'
BLUE = '0000FF'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
GROUP = re.compile(r'\[(\d+(?:\s*[,，;；–—-]\s*\d+)*)\]')


def element(tag, **attrs):
    result = OxmlElement(tag)
    for key, value in attrs.items():
        result.set(qn(key), str(value))
    return result


def visible_text(parent):
    return ''.join(t.text or '' for t in parent.iter(qn('w:t')))


def configure_citation_style(doc):
    style = doc.styles[STYLE] if STYLE in doc.styles else doc.styles.add_style(STYLE, WD_STYLE_TYPE.CHARACTER)
    style.base_style = None
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor.from_string(BLUE)
    style.font.superscript = True
    style.font.underline = False
    return style


def citation_run(run, doc):
    pr = run.get_or_add_rPr()
    values = [('w:rStyle', {'w:val': configure_citation_style(doc).style_id}),
              ('w:color', {'w:val': BLUE}), ('w:vertAlign', {'w:val': 'superscript'}),
              ('w:u', {'w:val': 'none'}), ('w:sz', {'w:val': 24}),
              ('w:rFonts', {'w:ascii': 'Times New Roman', 'w:hAnsi': 'Times New Roman', 'w:eastAsia': '宋体'})]
    for tag, attrs in values:
        old = pr.find(qn(tag))
        if old is not None:
            pr.remove(old)
        pr.append(element(tag, **attrs))


def text_run(text):
    run = element('w:r')
    value = element('w:t')
    value.set(qn('xml:space'), 'preserve')
    value.text = text
    run.append(value)
    return run


def add_field(parent, code, cached):
    runs = [element('w:r') for _ in range(5)]
    runs[0].append(element('w:fldChar', **{'w:fldCharType': 'begin'}))
    instruction = element('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' ' + code + ' '
    runs[1].append(instruction)
    runs[2].append(element('w:fldChar', **{'w:fldCharType': 'separate'}))
    runs[3].append(element('w:t'))
    runs[3][-1].text = str(cached)
    runs[4].append(element('w:fldChar', **{'w:fldCharType': 'end'}))
    parent.extend(runs)
    return runs


def bookmark_name(key):
    name = str(key) if str(key).startswith(PREFIX) else PREFIX + str(key)
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,39}', name):
        raise ValueError('Use an ASCII bibliography key with letters, numbers and underscores, at most 36 characters.')
    return name


def entry_nodes(doc, key, cached_number):
    name = bookmark_name(key)
    ids = [int(n.get(qn('w:id'))) for n in doc.element.iter(qn('w:bookmarkStart'))]
    bid = max(ids + [0]) + 1
    parent = element('w:p')
    parent.append(text_run('['))
    parent.append(element('w:bookmarkStart', **{'w:id': bid, 'w:name': name}))
    add_field(parent, f'SEQ {SEQUENCE} \\* ARABIC', cached_number)
    parent.append(element('w:bookmarkEnd', **{'w:id': bid}))
    parent.append(text_run('] '))
    return list(parent)


def add_bibliography_entry(doc, text, key, cached_number=None):
    """Append a reference whose bookmarked number is a native SEQ field."""
    if cached_number is None:
        cached_number = 1 + sum(n.get(qn('w:name'), '').startswith(PREFIX)
                                for n in doc.element.iter(qn('w:bookmarkStart')))
    name = bookmark_name(key)
    if any(n.get(qn('w:name')) == name for n in doc.element.iter(qn('w:bookmarkStart'))):
        raise ValueError('Bibliography key already exists: ' + name)
    paragraph = doc.add_paragraph(style='参考文献' if '参考文献' in doc.styles else 'Normal')
    paragraph._p.extend(entry_nodes(doc, key, cached_number))
    paragraph.add_run(text)
    return paragraph


def citation_nodes(doc, keys, numbers, separators=None):
    if len(keys) != len(numbers) or not keys:
        raise ValueError('Each citation key needs a cached number.')
    separators = separators if separators is not None else [','] * (len(keys) - 1)
    result = []
    for i, (key, number) in enumerate(zip(keys, numbers)):
        h = element('w:hyperlink', **{'w:anchor': bookmark_name(key), 'w:history': 1})
        if i == 0:
            h.append(text_run('['))
        add_field(h, f'REF {bookmark_name(key)} \\h \\* CHARFORMAT', number)
        if i == len(keys) - 1:
            h.append(text_run(']'))
        for run in h.iter(qn('w:r')):
            citation_run(run, doc)
        result.append(h)
        if i < len(keys) - 1:
            run = text_run(separators[i])
            citation_run(run, doc)
            result.append(run)
    return result


def add_citation(paragraph, keys, cached_numbers):
    """Append a clickable blue superscript [n] or [n,m] citation."""
    if isinstance(keys, (str, int)):
        keys, cached_numbers = [keys], [cached_numbers]
    nodes = citation_nodes(paragraph.part.document, keys, cached_numbers)
    paragraph._p.extend(nodes)
    return nodes


def _slice_runs(runs, start, end):
    result, offset = [], 0
    for run in runs:
        text = visible_text(run)
        a, b = max(start - offset, 0), min(end - offset, len(text))
        if a < b:
            copy = deepcopy(run)
            for child in list(copy):
                if child.tag != qn('w:rPr'):
                    copy.remove(child)
            value = element('w:t')
            value.set(qn('xml:space'), 'preserve')
            value.text = text[a:b]
            copy.append(value)
            result.append(copy)
        offset += len(text)
    return result


def _replace_text_runs(paragraph, pattern, replacement):
    """Match across adjacent text runs, preserving other inline objects and formatting."""
    groups, current = [], []
    for node in paragraph._p:
        plain = node.tag == qn('w:r') and all(c.tag in [qn('w:rPr'), qn('w:t')] for c in node)
        if plain:
            current.append(node)
        elif current:
            groups.append(current)
            current = []
    if current:
        groups.append(current)
    count = 0
    for runs in groups:
        text = ''.join(visible_text(r) for r in runs)
        replacements, cursor = [], 0
        for match in pattern.finditer(text):
            nodes = replacement(match)
            if nodes is None:
                continue
            replacements.extend(_slice_runs(runs, cursor, match.start()))
            replacements.extend(nodes)
            cursor = match.end()
            count += 1
        if cursor == 0:
            continue
        replacements.extend(_slice_runs(runs, cursor, len(text)))
        for node in replacements:
            runs[0].addprevious(node)
        for run in runs:
            paragraph._p.remove(run)
    return count


def _fields(parent):
    stack = []
    for node in parent.iter():
        if node.tag == qn('w:fldChar'):
            kind = node.get(qn('w:fldCharType'))
            if kind == 'begin':
                stack.append({'code': [], 'text': [], 'separate': False})
            elif kind == 'separate' and stack:
                stack[-1]['separate'] = True
            elif kind == 'end' and stack:
                yield stack.pop()
        elif node.tag == qn('w:instrText') and stack:
            stack[-1]['code'].append(node.text or '')
        elif node.tag == qn('w:t') and stack and stack[-1]['separate']:
            stack[-1]['text'].append(node)


def _set_cache(field, number):
    if not field['text']:
        raise ValueError('Field has no cached result: ' + ''.join(field['code']))
    field['text'][0].text = str(number)
    for node in field['text'][1:]:
        node.text = ''


def refresh_bibliography_cache(doc):
    """Refresh SEQ and REF cached results from bibliography order without opening Word."""
    numbers = {}
    for paragraph in doc.element.body.iter(qn('w:p')):
        names = [b.get(qn('w:name')) for b in paragraph.iter(qn('w:bookmarkStart'))
                 if b.get(qn('w:name'), '').startswith(PREFIX)]
        if not names:
            continue
        if len(names) != 1 or names[0] in numbers:
            raise ValueError('Duplicate or ambiguous bibliography bookmark.')
        numbers[names[0]] = len(numbers) + 1
        fields = [f for f in _fields(paragraph) if re.match(r'\s*SEQ\s+' + SEQUENCE + r'\b', ''.join(f['code']))]
        if len(fields) != 1:
            raise ValueError('Each bibliography entry needs one numbered SEQ field.')
        _set_cache(fields[0], numbers[names[0]])
    for field in _fields(doc.element.body):
        match = re.match(r'\s*REF\s+(Bib_\w+)\b', ''.join(field['code']))
        if match:
            if match[1] not in numbers:
                raise ValueError('Missing bibliography target: ' + match[1])
            _set_cache(field, numbers[match[1]])
    return numbers


def format_bibliography_fields(doc):
    configure_citation_style(doc)
    sid = doc.styles[STYLE].style_id
    for run in doc.element.body.iter(qn('w:r')):
        style = run.find('w:rPr/w:rStyle', NS)
        if style is not None and style.get(qn('w:val')) == sid:
            citation_run(run, doc)
    for link in doc.element.body.iter(qn('w:hyperlink')):
        if link.get(qn('w:anchor'), '').startswith(PREFIX):
            for run in link.iter(qn('w:r')):
                citation_run(run, doc)


def link_numeric_citations(doc):
    """Convert numbered reference paragraphs and matching plain-text citations in place."""
    paragraphs = [Paragraph(n, doc._body) for n in doc.element.body.iter(qn('w:p'))]
    records, lookup = [], {}
    for p in paragraphs:
        if p.style.name != '参考文献':
            continue
        match = re.match(r'^\[(\d+)\]\s+', visible_text(p._p))
        if not match:
            continue
        number = int(match[1])
        if number in lookup:
            raise ValueError('Ambiguous repeated reference number; process each bibliography scope separately.')
        existing = [b.get(qn('w:name')) for b in p._p.iter(qn('w:bookmarkStart'))
                    if b.get(qn('w:name'), '').startswith(PREFIX)]
        lookup[number] = existing[0] if existing else bookmark_name(number)
        records.append((p, number, bool(existing)))
    for p, number, existing in records:
        if not existing:
            _replace_text_runs(p, re.compile(r'^\[' + str(number) + r'\]\s+'),
                               lambda m, n=number: entry_nodes(doc, n, n))
    count = 0
    def replace(match):
        numbers = [int(n) for n in re.findall(r'\d+', match[1])]
        if not all(n in lookup for n in numbers):
            return None
        separators = re.split(r'\d+', match[1])[1:-1]
        return citation_nodes(doc, [lookup[n] for n in numbers], numbers, separators)
    for p in paragraphs:
        if p.style.name != '参考文献':
            count += _replace_text_runs(p, GROUP, replace)
    refresh_bibliography_cache(doc)
    format_bibliography_fields(doc)
    report = check_bibliography_fields(doc)
    report['converted_citation_groups'] = count
    return report


def check_bibliography_fields(doc):
    names = [b.get(qn('w:name')) for b in doc.element.body.iter(qn('w:bookmarkStart'))]
    managed = [n for n in names if n and n.startswith(PREFIX)]
    if len(managed) != len(set(managed)):
        raise ValueError('Bibliography bookmark names must be unique.')
    for run in doc.element.body.iter(qn('w:r')):
        style = run.find('w:rPr/w:rStyle', NS)
        if style is not None and style.get(qn('w:val')) == STYLE:
            color, align = run.find('w:rPr/w:color', NS), run.find('w:rPr/w:vertAlign', NS)
            if color is None or color.get(qn('w:val')) != BLUE or align is None or align.get(qn('w:val')) != 'superscript':
                raise ValueError('Citation punctuation must also be blue and superscript.')
    count = 0
    for link in doc.element.body.iter(qn('w:hyperlink')):
        target = link.get(qn('w:anchor'), '')
        if not target.startswith(PREFIX):
            continue
        if target not in managed:
            raise ValueError('Citation target does not exist: ' + target)
        fields = list(_fields(link))
        if len(fields) != 1 or not re.match(r'\s*REF\s+' + re.escape(target) + r'\s', ''.join(fields[0]['code'])):
            raise ValueError('Citation hyperlink must contain its matching native REF field.')
        code = ''.join(fields[0]['code'])
        if not re.search(r'\\h\b', code) or not re.search(r'\\\*\s+CHARFORMAT\b', code):
            raise ValueError('Citation REF field needs hyperlink and character-format switches.')
        for run in link.iter(qn('w:r')):
            color, align = run.find('w:rPr/w:color', NS), run.find('w:rPr/w:vertAlign', NS)
            if color is None or color.get(qn('w:val')) != BLUE or align is None or align.get(qn('w:val')) != 'superscript':
                raise ValueError('Citation must be blue and superscript, including brackets.')
        count += 1
    return {'bibliography_entries': len(managed), 'linked_reference_fields': count}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args()
    doc = Document(args.input)
    print(json.dumps(link_numeric_citations(doc), ensure_ascii=False))
    doc.save(args.out)
