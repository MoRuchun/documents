"""Apply the user's DOCX styles and provide reusable authoring helpers."""
from __future__ import annotations
import argparse
from copy import deepcopy
from pathlib import Path
import json
import re
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from bibliography_fields import format_bibliography_fields, check_bibliography_fields

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
BLUE = '0000FF'
NUMBERING_NAME = 'UserDocumentHeadings'
PUNCTUATION = set('，。；：！？、（）【】《》“”‘’—…')


def node(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_child(parent, tag, **attrs):
    old = parent.find(qn(tag))
    if old is None:
        old = node(tag)
        parent.append(old)
    old.attrib.clear()
    for k, v in attrs.items():
        old.set(qn(k), str(v))
    return old


def remove_children(parent, *tags):
    for tag in tags:
        for e in list(parent.findall(qn(tag))):
            parent.remove(e)


def set_fonts(rpr, east='宋体', western='Times New Roman'):
    f = rpr.get_or_add_rFonts()
    f.attrib.clear()
    for key, value in [('ascii', western), ('hAnsi', western), ('eastAsia', east), ('cs', western)]:
        f.set(qn('w:' + key), value)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = node('w:lang'); rpr.append(lang)
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'zh-CN')


# name: East Asian font, size, bold, color, alignment, before, after,
# line height in pt (or multiple), multiple flag, first-line characters, left characters.
PARAGRAPH_STYLES = {
    'Normal': ('宋体',12,False,'000000','justify',0,0,22,False,2,0),
    'Heading 1': ('黑体',16,True,'000000','left',0,6,24,False,0,0),
    'Heading 2': ('黑体',14,True,'000000','left',6,6,22,False,0,0),
    'Heading 3': ('黑体',12,True,'000000','left',6,6,20,False,0,0),
    '无编号标题': ('黑体',12,True,'000000','left',6,6,20,False,0,0),
    '分项标题': ('楷体',12,True,'110CE6','left',3,3,20,False,0,0),
    '图片段落': ('宋体',12,False,'000000','center',0,0,1,True,0,0),
    '图题': ('黑体',9,False,'000000','center',0,6,1,True,0,0),
    '表题': ('黑体',9,False,'000000','center',6,0,1,True,0,0),
    '表内文字': ('宋体',10.5,False,'000000','left',0,0,1,True,0,0),
    '图表注释': ('宋体',9,False,'000000','left',0,0,1,True,0,0),
    '行间公式': ('宋体',12,False,'000000','left',3,3,1,True,0,0),
    '参考文献': ('宋体',10.5,False,'000000','justify',0,0,18,False,0,0),
    '页码': ('宋体',9,False,'000000','center',0,0,1,True,0,0),
    '罗马页码': ('宋体',9,False,'000000','center',0,0,1,True,0,0),
    'TOC 1': ('黑体',12,False,'000000','left',0,0,1.5,True,0,0),
    'TOC 2': ('宋体',12,False,'000000','left',0,0,1.5,True,0,2),
    'TOC 3': ('宋体',12,False,'000000','left',0,0,1.5,True,0,4),
}


def configure_styles(doc):
    """Create the complete paragraph, character and table style set."""
    for name, values in PARAGRAPH_STYLES.items():
        east,size,bold,color,align,before,after,line,multiple,first,left = values
        st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != 'Normal':
            st.base_style = None
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor.from_string(color)
        set_fonts(st.element.get_or_add_rPr(), east, '宋体' if name == '罗马页码' else 'Times New Roman')
        f = st.paragraph_format
        f.alignment = {'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        f.space_before, f.space_after = Pt(before), Pt(after)
        f.left_indent, f.right_indent, f.first_line_indent = Pt(left * size), Pt(0), Pt(first * size)
        f.keep_with_next = name == 'Heading 1'
        f.keep_together = False
        f.page_break_before = name == 'Heading 1'
        f.widow_control = True
        pp = st.element.get_or_add_pPr()
        spacing = pp.get_or_add_spacing()
        spacing.attrib.clear()
        for key,value in [('before',round(before*20)),('after',round(after*20)),
                          ('line',round(line*240) if multiple else round(line*20)),
                          ('lineRule','auto' if multiple else 'exact')]:
            spacing.set(qn('w:'+key),str(value))
        if name == '参考文献':
            spacing.set(qn('w:afterLines'),'20')
        ind = pp.get_or_add_ind()
        ind.attrib.clear()
        for key,value in [('left',round(left*size*20)),('leftChars',left*100),('right',0),('firstLine',round(first*size*20)),('firstLineChars',first*100)]:
            ind.set(qn('w:'+key),str(value))
        if name == '参考文献':
            ind.attrib.clear()
            for key,value in [('left',0),('leftChars',0),('right',0),('hanging',420),('hangingChars',200)]:
                ind.set(qn('w:'+key),str(value))
        set_child(pp,'w:snapToGrid',**{'w:val':'0'})
        set_child(pp,'w:outlineLvl',**{'w:val':int(name[-1])-1 if name.startswith('Heading ') else 9})
        if name.startswith('TOC '):
            remove_children(pp,'w:tabs')
            f.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if name == '行间公式':
            remove_children(pp,'w:tabs')
            f.tab_stops.add_tab_stop(Cm(7.75), WD_TAB_ALIGNMENT.CENTER)
            f.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT)
    for name,east,west,color in [('交叉引用','宋体','Times New Roman',BLUE),('数学字符','Cambria Math','Cambria Math','000000')]:
        st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.CHARACTER)
        st.base_style = None
        set_fonts(st.element.get_or_add_rPr(),east,west)
        st.font.color.rgb = RGBColor.from_string(color)
    st = doc.styles['三线表'] if '三线表' in doc.styles else doc.styles.add_style('三线表',WD_STYLE_TYPE.TABLE)
    st.base_style = None
    st.font.size = Pt(10.5)
    st.font.bold = False
    set_fonts(st.element.get_or_add_rPr())
    pp = st.element.get_or_add_pPr()
    set_child(pp,'w:spacing',**{'w:before':0,'w:after':0,'w:line':240,'w:lineRule':'auto'})
    set_child(pp,'w:snapToGrid',**{'w:val':'0'})
    for st in doc.styles:
        if st.name != 'Normal':
            st.base_style = None


def configure_page(doc):
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin,sec.right_margin,sec.gutter = Cm(3),Cm(2.5),Cm(0)
        remove_children(sec._sectPr,'w:docGrid','w:headerReference')


def configure_heading_numbering(doc):
    """One stable numbering definition, linked both to styles and paragraphs."""
    root = doc.part.numbering_part.element
    abstract = next((e for e in root.findall(qn('w:abstractNum'))
                     if e.find(qn('w:name')) is not None and e.find(qn('w:name')).get(qn('w:val')) == NUMBERING_NAME), None)
    if abstract is None:
        aid=max([int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))]+[-1])+1
        abstract=node('w:abstractNum',**{'w:abstractNumId':aid})
        first_num=root.find(qn('w:num'))
        if first_num is None:root.append(abstract)
        else:root.insert(root.index(first_num),abstract)
    else:
        aid=int(abstract.get(qn('w:abstractNumId')))
        abstract.clear();abstract.set(qn('w:abstractNumId'),str(aid))
    abstract.append(node('w:multiLevelType',**{'w:val':'multilevel'}))
    abstract.append(node('w:name',**{'w:val':NUMBERING_NAME}))
    for level in range(3):
        lv=node('w:lvl',**{'w:ilvl':level})
        lv.append(node('w:start',**{'w:val':1}))
        lv.append(node('w:numFmt',**{'w:val':'decimal'}))
        lv.append(node('w:pStyle',**{'w:val':doc.styles[f'Heading {level+1}'].style_id}))
        lv.append(node('w:suff',**{'w:val':'space'}))
        lv.append(node('w:lvlText',**{'w:val':'.'.join('%'+str(i+1) for i in range(level+1))}))
        lv.append(node('w:lvlJc',**{'w:val':'left'}))
        pp=node('w:pPr');pp.append(node('w:ind',**{'w:left':0,'w:hanging':0}));lv.append(pp)
        # Correctly typed rPr keeps font accessors and schema ordering available.
        rr=OxmlElement('w:rPr');set_fonts(rr,'黑体');rr.get_or_add_b().val=True;lv.append(rr)
        abstract.append(lv)
    num=next((e for e in root.findall(qn('w:num')) if e.find(qn('w:abstractNumId')) is not None and e.find(qn('w:abstractNumId')).get(qn('w:val'))==str(aid)),None)
    if num is None:
        nid=max([int(e.get(qn('w:numId'))) for e in root.findall(qn('w:num'))]+[0])+1
        num=node('w:num',**{'w:numId':nid});num.append(node('w:abstractNumId',**{'w:val':aid}));root.append(num)
    nid=int(num.get(qn('w:numId')))
    # WordprocessingML orders abstractNum elements before num elements.
    for abstract_node in list(root.findall(qn('w:abstractNum'))):
        root.remove(abstract_node)
        pos=next((i for i,e in enumerate(root) if e.tag==qn('w:num')),len(root))
        root.insert(pos,abstract_node)
    for level in range(3):
        pp=doc.styles[f'Heading {level+1}'].element.get_or_add_pPr()
        set_numpr(pp,nid,level)
    return nid


def set_numpr(pp,nid,level):
    np=pp.get_or_add_numPr()
    np.get_or_add_ilvl().val=level
    np.get_or_add_numId().val=nid


def apply_paragraph_style(p,doc,name,num_id=None):
    p.style=doc.styles[name]
    pp=p._p.get_or_add_pPr()
    source=doc.styles[name].element.get_or_add_pPr()
    # Replace layout overrides while retaining bookmarks, fields and section boundaries.
    for tag in ['w:spacing','w:ind','w:jc','w:keepNext','w:keepLines','w:pageBreakBefore','w:widowControl','w:snapToGrid','w:outlineLvl']:
        remove_children(pp,tag)
        value=source.find(qn(tag))
        if value is not None:pp.append(deepcopy(value))
    if name.startswith('Heading '):
        set_numpr(pp,num_id,int(name[-1])-1)
    elif name in ['无编号标题','分项标题']:
        remove_children(pp,'w:numPr')
    if name == '行间公式':
        remove_children(pp,'w:tabs')
        pp.append(deepcopy(source.find(qn('w:tabs'))))
    # Strip only overrides governed by this specification; preserve body emphasis.
    east,size,bold,color,*_=PARAGRAPH_STYLES[name]
    for r in p._p.findall(qn('w:r')):
        rp=r.get_or_add_rPr()
        rstyle=rp.find(qn('w:rStyle'))
        if rstyle is not None and rstyle.get(qn('w:val'))==doc.styles['交叉引用'].style_id:
            continue
        set_fonts(rp,east,'宋体' if name=='罗马页码' else 'Times New Roman')
        rp.get_or_add_sz().val=Pt(size)
        rp.get_or_add_color().val=RGBColor.from_string(color)
        if name!='Normal':rp.get_or_add_b().val=bold
    color_chinese_punctuation(p)


def color_chinese_punctuation(p):
    for r in list(p._p.findall(qn('w:r'))):
        t=r.find(qn('w:t'))
        if t is None or not t.text or not any(c in PUNCTUATION for c in t.text):continue
        if any(r.find(qn(tag)) is not None for tag in ['w:fldChar','w:instrText','w:drawing']):continue
        if len(r.findall(qn('w:t')))!=1:continue
        chunks=re.findall(r'[，。；：！？、（）【】《》“”‘’—…]+|[^，。；：！？、（）【】《》“”‘’—…]+',t.text)
        for text in chunks:
            rr=deepcopy(r);rr.find(qn('w:t')).text=text
            if text[0] in PUNCTUATION:set_fonts(rr.get_or_add_rPr(),'宋体','宋体')
            r.addprevious(rr)
        p._p.remove(r)


def format_tables(doc,num_id):
    for tbl_node in doc.element.body.iter(qn('w:tbl')):
        pr=tbl_node.find(qn('w:tblPr'))
        if pr is None:pr=node('w:tblPr');tbl_node.insert(0,pr)
        set_child(pr,'w:tblStyle',**{'w:val':doc.styles['三线表'].style_id})
        set_child(pr,'w:tblW',**{'w:type':'dxa','w:w':round(Cm(15.5).twips)})
        set_child(pr,'w:jc',**{'w:val':'center'})
        set_child(pr,'w:tblLayout',**{'w:type':'fixed'})
        borders=set_child(pr,'w:tblBorders');borders.clear()
        for edge in ['top','left','bottom','right','insideH','insideV']:
            borders.append(node('w:'+edge,**{'w:val':'nil'}))
        margin=set_child(pr,'w:tblCellMar');margin.clear()
        for edge,val in [('top',.1),('bottom',.1),('left',.12),('right',.12)]:
            margin.append(node('w:'+edge,**{'w:w':round(Cm(val).twips),'w:type':'dxa'}))
        grid=tbl_node.find(qn('w:tblGrid'))
        cols=grid.findall(qn('w:gridCol')) if grid is not None else []
        widths=[int(c.get(qn('w:w'),'0')) for c in cols]
        if widths and sum(widths)>0:
            total=round(Cm(15.5).twips)
            updated=[round(x/sum(widths)*total) for x in widths]
            updated[-1]+=total-sum(updated)
            for c,width in zip(cols,updated):c.set(qn('w:w'),str(width))
        rows=tbl_node.findall(qn('w:tr'))
        for ri,row in enumerate(rows):
            trpr=row.get_or_add_trPr();set_child(trpr,'w:cantSplit',**{'w:val':1})
            if ri==0:set_child(trpr,'w:tblHeader',**{'w:val':1})
            colidx=0
            for cell in row.findall(qn('w:tc')):
                cp=cell.get_or_add_tcPr();set_child(cp,'w:vAlign',**{'w:val':'center'})
                remove_children(cp,'w:shd','w:tcMar')
                span=cp.find(qn('w:gridSpan'));n=int(span.get(qn('w:val'))) if span is not None else 1
                if widths and sum(widths)>0:
                    set_child(cp,'w:tcW',**{'w:type':'dxa','w:w':sum(updated[colidx:colidx+n])})
                colidx+=n
                b=set_child(cp,'w:tcBorders');b.clear()
                for edge in ['top','left','bottom','right']:
                    line=(ri==0 and edge in ['top','bottom']) or (ri==len(rows)-1 and edge=='bottom')
                    size=6 if ri==0 and edge=='bottom' else 12
                    b.append(node('w:'+edge,**{'w:val':'single' if line else 'nil','w:sz':size,'w:color':'000000'}))
                for pnode in cell.findall(qn('w:p')):
                    apply_paragraph_style(Paragraph(pnode,doc._body),doc,'表内文字',num_id)


def format_cross_references(doc):
    sid=doc.styles['交叉引用'].style_id
    def blue(r):
        rp=r.get_or_add_rPr();set_child(rp,'w:rStyle',**{'w:val':sid})
        set_fonts(rp);rp.get_or_add_color().val=RGBColor.from_string(BLUE)
    for hyperlink in doc.element.body.iter(qn('w:hyperlink')):
        instruction=' '.join(t.text or '' for t in hyperlink.iter(qn('w:instrText')))
        if not re.search(r'\bREF\s',instruction) or re.search(r'\bPAGEREF\s',instruction):continue
        for r in hyperlink.iter(qn('w:r')):blue(r)
    # Also handle standard REF fields outside hyperlink wrapper elements.
    stack=[]
    for e in doc.element.body.iter():
        if e.tag==qn('w:fldChar'):
            typ=e.get(qn('w:fldCharType'))
            if typ=='begin':stack.append({'code':[],'runs':[]})
            elif typ=='end' and stack:
                ctx=stack.pop();code=''.join(x.text or '' for x in ctx['code'])
                if re.match(r'\s*REF\s+',code):
                    # REF in a table caption supplies the continued table number;
                    # only prose cross-references use the blue character style.
                    paragraph=e.getparent()
                    while paragraph is not None and paragraph.tag!=qn('w:p'):paragraph=paragraph.getparent()
                    ps=paragraph.find('w:pPr/w:pStyle',NS) if paragraph is not None else None
                    styleid=ps.get(qn('w:val')) if ps is not None else ''
                    if styleid in [doc.styles['表题'].style_id,doc.styles['图题'].style_id]:continue
                    if not re.search(r'\\h\b',code):ctx['code'][-1].text=(ctx['code'][-1].text or '').rstrip()+' \\h '
                    if not re.search(r'\\\*\s+CHARFORMAT\b',code):ctx['code'][-1].text=(ctx['code'][-1].text or '').rstrip()+' \\* CHARFORMAT '
                    for r in ctx['runs']:blue(r)
        elif e.tag==qn('w:instrText') and stack:stack[-1]['code'].append(e)
        elif e.tag==qn('w:r'):
            for ctx in stack:ctx['runs'].append(e)


def format_math(doc):
    for r in doc.element.body.iter(qn('m:r')):
        pr=r.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
        set_fonts(pr,'Cambria Math','Cambria Math')
    mp=doc.settings.element.find(qn('m:mathPr'))
    if mp is None:mp=node('m:mathPr');doc.settings.element.append(mp)
    set_child(mp,'m:mathFont',**{'m:val':'Cambria Math'})


def apply_document(doc,style_map=None):
    """Batch-format an already structured document without recreating its content."""
    configure_styles(doc);configure_page(doc)
    nid=configure_heading_numbering(doc)
    style_map=style_map or {}
    id_names={st.style_id:st.name for st in doc.styles}
    for pnode in doc.element.body.iter(qn('w:p')):
        if any(a.tag==qn('w:tc') for a in pnode.iterancestors()):continue
        p=Paragraph(pnode,doc._body)
        current=p.style.name
        name=style_map.get(current,style_map.get(p.style.style_id,current))
        if name in PARAGRAPH_STYLES:apply_paragraph_style(p,doc,name,nid)
        else:set_child(pnode.get_or_add_pPr(),'w:snapToGrid',**{'w:val':0})
    format_tables(doc,nid);format_cross_references(doc);format_bibliography_fields(doc);format_math(doc)
    for sec in doc.sections:
        for footer in [sec.footer,sec.first_page_footer,sec.even_page_footer]:
            for p in footer.paragraphs:
                if not (p.text or p._p.findall('.//'+qn('w:instrText'))):continue
                name='罗马页码' if p.style.name=='罗马页码' else '页码'
                apply_paragraph_style(p,doc,name,nid)
    set_child(doc.settings.element,'w:updateFields',**{'w:val':'false'})
    return doc


def new_document():
    return apply_document(Document())


def add_heading(doc,text,level=1):
    if level not in [1,2,3]:raise ValueError('Heading level must be 1, 2 or 3')
    nid=configure_heading_numbering(doc)
    p=doc.add_paragraph(text,f'Heading {level}')
    apply_paragraph_style(p,doc,f'Heading {level}',nid)
    return p


def add_field(parent,code,result='',blue=False):
    def run():
        r=node('w:r');parent.append(r)
        if blue:
            pr=r.get_or_add_rPr();set_child(pr,'w:rStyle',**{'w:val':'交叉引用'})
            set_fonts(pr);pr.get_or_add_color().val=RGBColor.from_string(BLUE)
        return r
    if blue and not re.search(r'\\\*\s+CHARFORMAT',code):code+=' \\* CHARFORMAT'
    run().append(node('w:fldChar',**{'w:fldCharType':'begin'}))
    instr=node('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' '+code+' ';run().append(instr)
    run().append(node('w:fldChar',**{'w:fldCharType':'separate'}))
    t=node('w:t');t.set(qn('xml:space'),'preserve');t.text=result;run().append(t)
    run().append(node('w:fldChar',**{'w:fldCharType':'end'}))


def add_cross_reference(paragraph,bookmark,cached_number,prefix=''):
    h=node('w:hyperlink',**{'w:anchor':bookmark,'w:history':1});paragraph._p.append(h)
    if prefix:
        r=node('w:r');pr=r.get_or_add_rPr();set_fonts(pr)
        set_child(pr,'w:rStyle',**{'w:val':'交叉引用'});pr.get_or_add_color().val=RGBColor.from_string(BLUE)
        t=node('w:t');t.set(qn('xml:space'),'preserve');t.text=prefix;r.append(t);h.append(r)
    add_field(h,f'REF {bookmark} \\h',cached_number,blue=True)
    return h



def add_bookmark(parent, name):
    root=parent.getroottree().getroot()
    ids=[int(e.get(qn('w:id'))) for e in root.iter(qn('w:bookmarkStart'))]
    bid=max(ids+[0])+1
    start=node('w:bookmarkStart',**{'w:id':bid,'w:name':name});parent.append(start)
    return bid


def end_bookmark(parent, bid):
    parent.append(node('w:bookmarkEnd',**{'w:id':bid}))


def add_caption(doc, kind, title, bookmark, chapter, index, continued=False):
    """Add a chapter-based SEQ caption, or a continued table caption referencing it."""
    if kind not in ['图','表']:raise ValueError('Caption kind must be 图 or 表')
    p=doc.add_paragraph(style='图题' if kind=='图' else '表题')
    if continued:
        if kind!='表':raise ValueError('Only table captions have a continued form')
        p.add_run('续表 ')
        add_field(p._p,f'REF {bookmark} \\h',f'{chapter}-{index}')
    else:
        p.add_run(kind+' ')
        bid=add_bookmark(p._p,bookmark)
        add_field(p._p,'STYLEREF 1 \\s',str(chapter));p.add_run('-')
        add_field(p._p,f'SEQ {"Figure" if kind=="图" else "Table"} \\* ARABIC \\s 1',str(index))
        end_bookmark(p._p,bid)
    p.add_run(' '+title)
    apply_paragraph_style(p,doc,'图题' if kind=='图' else '表题',None)
    return p


def add_table(doc, headers, rows, widths_cm=None):
    """Create a full-width three-line table with explicit single-spaced cells."""
    table=doc.add_table(rows=1,cols=len(headers))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=widths_cm or [15.5/len(headers)]*len(headers)
    total=sum(widths)
    widths=[w/total*15.5 for w in widths]
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    for i,values in enumerate([headers]+list(rows)):
        row=table.rows[0] if i==0 else table.add_row()
        for cell,text,width in zip(row.cells,values,widths):
            cell.width=Cm(width)
            cell.paragraphs[0].style=doc.styles['表内文字']
            cell.paragraphs[0].add_run(str(text))
    format_tables(doc,configure_heading_numbering(doc))
    return table


def add_inline_math(paragraph, elements):
    obj=node('m:oMath')
    for e in elements:obj.append(deepcopy(e))
    paragraph._p.append(obj)
    return obj


def add_numbered_equation(doc, elements, bookmark, chapter, index):
    p=doc.add_paragraph(style='行间公式')
    apply_paragraph_style(p,doc,'行间公式')
    p.add_run('\t');add_inline_math(p,elements);p.add_run('\t')
    bid=add_bookmark(p._p,bookmark)
    p.add_run('（');add_field(p._p,'STYLEREF 1 \\s',str(chapter));p.add_run('-')
    add_field(p._p,'SEQ Equation \\* ARABIC \\s 1',str(index));p.add_run('）')
    end_bookmark(p._p,bid)
    return p

def check_document(doc):
    """Check the operational invariants affected by batch formatting."""
    errors=[]
    root=doc.element.body
    nums=doc.part.numbering_part.element
    heading_ids={doc.styles[f'Heading {i}'].style_id:i-1 for i in range(1,4)}
    labels=[];counters=[0,0,0]
    for p in root.iter(qn('w:p')):
        style=p.find('w:pPr/w:pStyle',NS)
        if style is None or style.get(qn('w:val')) not in heading_ids:continue
        level=heading_ids[style.get(qn('w:val'))]
        numpr=p.find('w:pPr/w:numPr',NS)
        if numpr is None or numpr.find(qn('w:ilvl')) is None or int(numpr.find(qn('w:ilvl')).get(qn('w:val')))!=level:
            errors.append('Heading paragraph lacks its direct numbering level');continue
        n=numpr.find(qn('w:numId'));nid=n.get(qn('w:val')) if n is not None else None
        num=next((e for e in nums.findall(qn('w:num')) if e.get(qn('w:numId'))==nid),None)
        if num is None:errors.append('Heading numbering definition missing');continue
        aid=num.find(qn('w:abstractNumId')).get(qn('w:val'))
        abstract=next((e for e in nums.findall(qn('w:abstractNum')) if e.get(qn('w:abstractNumId'))==aid),None)
        lv=next((e for e in abstract.findall(qn('w:lvl')) if int(e.get(qn('w:ilvl')))==level),None) if abstract is not None else None
        expect='.'.join('%'+str(i+1) for i in range(level+1))
        if lv is None or lv.find(qn('w:lvlText')).get(qn('w:val'))!=expect:errors.append('Heading numbering pattern incorrect')
        if lv is not None and lv.find(qn('w:lvlRestart')) is not None:errors.append('Heading restart must use the Word-compatible default')
        counters[level]+=1
        for i in range(level+1,3):counters[i]=0
        labels.append('.'.join(map(str,counters[:level+1])))
    table_ps=list(root.findall('.//w:tc/w:p',NS))
    for p in table_ps:
        sp=p.find('w:pPr/w:spacing',NS)
        if sp is None or sp.get(qn('w:line'))!='240' or sp.get(qn('w:lineRule'))!='auto':errors.append('Table paragraph is not single spaced')
    for name in ['参考文献']:
        ind=doc.styles[name].element.find('w:pPr/w:ind',NS)
        if ind is None or ind.get(qn('w:left'))!='0' or ind.get(qn('w:leftChars'))!='0' or ind.get(qn('w:hangingChars'))!='200':errors.append('Reference indentation incorrect')
    ref_run_count=0
    for h in root.iter(qn('w:hyperlink')):
        code=' '.join(t.text or '' for t in h.iter(qn('w:instrText')))
        if not re.search(r'\bREF\s',code) or re.search(r'\bPAGEREF\s',code):continue
        for r in h.iter(qn('w:r')):
            col=r.find('w:rPr/w:color',NS)
            if col is None or col.get(qn('w:val'))!=BLUE:errors.append('Cross-reference run lacks explicit blue')
            if r.find(qn('w:t')) is not None:ref_run_count+=1
        if not re.search(r'\\\*\s+CHARFORMAT',code):errors.append('Cross-reference field lacks CHARFORMAT')
    if errors:raise ValueError('\n'.join(sorted(set(errors))))
    bibliography_report=check_bibliography_fields(doc)
    return {'heading_numbers':labels,'table_paragraphs_single_spaced':len(table_ps),
            'cross_reference_text_runs_blue':ref_run_count,'reference_left_indent':0,
            'reference_hanging_characters':2,'native_equations':len(list(root.iter(qn('m:oMath')))),
            **bibliography_report}


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('input',type=Path)
    ap.add_argument('--out',type=Path,help='Output DOCX; required except for --check-only')
    ap.add_argument('--style-map',type=Path,help='JSON object mapping input style names to user style names')
    ap.add_argument('--check-only',action='store_true')
    args=ap.parse_args()
    doc=Document(args.input)
    if not args.check_only:
        if args.out is None:ap.error('--out is required')
        mapping=json.loads(args.style_map.read_text()) if args.style_map else None
        apply_document(doc,mapping)
        report=check_document(doc)
        args.out.parent.mkdir(parents=True,exist_ok=True);doc.save(args.out)
    else:report=check_document(doc)
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
