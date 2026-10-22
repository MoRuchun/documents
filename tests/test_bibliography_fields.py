import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import user_formatting as fmt
import bibliography_fields as bib


class BibliographyTests(unittest.TestCase):
    def test_split_runs_grouped_repeated_and_round_trip(self):
        doc = fmt.new_document()
        p = doc.add_paragraph('研究', 'Normal')
        p.add_run('表明').bold = True
        p.add_run('[1,')
        p.add_run('2]，又见[1]。')
        doc.add_paragraph('[1] 文献甲', '参考文献')
        doc.add_paragraph('[2] 文献乙', '参考文献')
        text = bib.visible_text(p._p)
        report = bib.link_numeric_citations(doc)
        self.assertEqual(report['linked_reference_fields'], 3)
        self.assertEqual(bib.visible_text(p._p), text)
        self.assertTrue(next(r for r in p.runs if r.text == '表明').bold)
        fmt.apply_document(doc)
        self.assertEqual(bib.link_numeric_citations(doc)['converted_citation_groups'], 0)
        with tempfile.TemporaryDirectory() as td:
            target = Path(td) / 'citations.docx'
            doc.save(target)
            reopened = Document(target)
            self.assertEqual(bib.check_bibliography_fields(reopened)['linked_reference_fields'], 3)
            self.assertEqual(fmt.check_document(reopened)['bibliography_entries'], 2)

    def test_reordering_refreshes_entry_and_citation_numbers(self):
        doc = fmt.new_document()
        p = doc.add_paragraph('结果', 'Normal')
        bib.add_citation(p, ['alpha', 'beta'], [1, 2])
        a = bib.add_bibliography_entry(doc, '甲', 'alpha', 1)
        b = bib.add_bibliography_entry(doc, '乙', 'beta', 2)
        a._p.addprevious(b._p)
        bib.refresh_bibliography_cache(doc)
        self.assertEqual(bib.visible_text(p._p), '结果[2,1]')
        self.assertEqual(bib.visible_text(a._p), '[2] 甲')
        self.assertEqual(bib.visible_text(b._p), '[1] 乙')
        fmt.apply_document(doc)
        for r in p._p.iter(qn('w:r')):
            if r.find('w:rPr/w:rStyle', bib.NS) is not None:
                self.assertEqual(r.find('w:rPr/w:color', bib.NS).get(qn('w:val')), '0000FF')
                self.assertEqual(r.find('w:rPr/w:vertAlign', bib.NS).get(qn('w:val')), 'superscript')

    def test_dangling_targets_and_duplicate_entry_numbers(self):
        doc = fmt.new_document()
        p = doc.add_paragraph('见', 'Normal')
        bib.add_citation(p, 'missing', 1)
        with self.assertRaisesRegex(ValueError, 'does not exist'):
            bib.check_bibliography_fields(doc)
        doc = fmt.new_document()
        doc.add_paragraph('[1] 甲', '参考文献')
        doc.add_paragraph('[1] 乙', '参考文献')
        with self.assertRaisesRegex(ValueError, 'Ambiguous'):
            bib.link_numeric_citations(doc)


if __name__ == '__main__':
    unittest.main()
