import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import user_formatting as fmt

class FormattingTests(unittest.TestCase):
    def test_numbering_reset_and_round_trip(self):
        doc=fmt.new_document()
        for level in [1,2,3,2,1,2]:fmt.add_heading(doc,'标题',level)
        fmt.apply_document(doc)
        n=len(doc.part.numbering_part.element.findall(qn('w:abstractNum')))
        fmt.apply_document(doc)
        self.assertEqual(len(doc.part.numbering_part.element.findall(qn('w:abstractNum'))),n)
        self.assertEqual(fmt.check_document(doc)['heading_numbers'],['1','1.1','1.1.1','1.2','2','2.1'])
        custom=[a for a in doc.part.numbering_part.element.findall(qn('w:abstractNum')) if a.find(qn('w:name')) is not None and a.find(qn('w:name')).get(qn('w:val'))==fmt.NUMBERING_NAME]
        self.assertEqual(len(custom),1)
        self.assertFalse(custom[0].findall('.//'+qn('w:lvlRestart')))
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'sample.docx';doc.save(p)
            self.assertEqual(fmt.check_document(Document(p)),fmt.check_document(doc))

    def test_complete_reference_color(self):
        doc=fmt.new_document();fmt.add_heading(doc,'记录',1)
        fmt.add_caption(doc,'表','地块记录','records',1,1)
        p=doc.add_paragraph('见','Normal');h=fmt.add_cross_reference(p,'records','1-1',prefix='表 ')
        for r in h.iter(qn('w:r')):r.get_or_add_rPr().get_or_add_color().val=RGBColor(0,0,0)
        fmt.apply_document(doc)
        self.assertEqual(''.join(t.text or '' for t in h.iter(qn('w:t'))),'表 1-1')
        for r in h.iter(qn('w:r')):self.assertEqual(r.find('w:rPr/w:color',fmt.NS).get(qn('w:val')),'0000FF')
        self.assertIn('CHARFORMAT',''.join(t.text or '' for t in h.iter(qn('w:instrText'))))

    def test_table_overrides(self):
        doc=fmt.new_document();t=fmt.add_table(doc,['地块','温度'],[['A','28.5']],[8,7.5])
        for row in t.rows:
            for cell in row.cells:cell.paragraphs[0].paragraph_format.line_spacing=Pt(35)
        fmt.apply_document(doc)
        for row in t.rows:
            for cell in row.cells:
                sp=cell.paragraphs[0]._p.find('w:pPr/w:spacing',fmt.NS)
                self.assertEqual((sp.get(qn('w:line')),sp.get(qn('w:lineRule'))),('240','auto'))

    def test_reference_indent_and_no_opening_prompt(self):
        doc=fmt.new_document();p=doc.add_paragraph('[1] 示例条目','参考文献');p.paragraph_format.left_indent=Pt(24)
        fmt.apply_document(doc)
        ind=p._p.find('w:pPr/w:ind',fmt.NS)
        self.assertEqual((ind.get(qn('w:left')),ind.get(qn('w:leftChars')),ind.get(qn('w:hangingChars'))),('0','0','200'))
        self.assertEqual(doc.settings.element.find(qn('w:updateFields')).get(qn('w:val')),'false')
        self.assertEqual(p.text,'[1] 示例条目')

if __name__=='__main__':unittest.main()
